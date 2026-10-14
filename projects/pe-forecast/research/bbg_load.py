"""使用者提供的 Bloomberg 匯出檔（.xlsx，不進 repo）：518 檔、2019-01～2026-09 月資料。

工作表：Static_Info（GICS）、1_Price 月底收盤價（已分割還原）、2_EPS_1BF（未來 12 個月共識 EPS）、3_EPS_2BF（第 13–24 個月）、
4_EPS_3BF（第 25–36 個月）、5_ROIC、6_Margin（季）、7_MktCap（月底市值）。
匯出時每一列都靠左排（每檔從自己第一筆資料開始放在第一欄、中間沒有空格），只有股價表的表頭有 2019-01-31 和 2026-09-30 兩個日期
→ 讀進來時把每一列靠右對齊到 2026-09（最後一筆＝2026-09-30）。對齊方式用面板股價驗證（research/bbg_consensus_test.py 第 0 節）。
"""
import numpy as np
import pandas as pd

SHEETS = {"px": "1_Price", "f1": "2_EPS_1BF", "f2": "3_EPS_2BF", "f3": "4_EPS_3BF", "mc": "7_MktCap"}
END = pd.Timestamp("2026-09-01")
N = 93                                                   # 2019-01 … 2026-09


def load(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    names = {ws.title: ws for ws in wb.worksheets}
    months = pd.date_range(END - pd.DateOffset(months=N - 1), END, freq="MS")
    out = {}
    for key, pre in SHEETS.items():
        ws = next(v for k, v in names.items() if k.startswith(pre))
        rows = list(ws.iter_rows(values_only=True))[1:]
        mat = {}
        for r in rows:
            if not isinstance(r[0], str):
                continue
            t = r[0].replace("US Equity", "").strip().replace(" ", "-").replace("/", "-")
            v = [float(c) if isinstance(c, (int, float)) else np.nan for c in r[1:N + 1]]
            k = int(np.sum(np.isfinite(v)))
            vals = np.full(N, np.nan)
            if k:
                vals[N - k:] = np.array(v[:k], float)    # 靠右對齊：最後一筆＝2026-09
            mat[t] = vals
        out[key] = pd.DataFrame(mat, index=months)
    st = next(v for k, v in names.items() if k.startswith("Static"))
    rows = list(st.iter_rows(values_only=True))
    info = pd.DataFrame([[str(r[0]).strip()] + list(r[1:5]) for r in rows[1:] if r[0]], columns=["ticker", "gics1", "gics2", "gics4", "adv30"])
    return out, info
